#!/usr/bin/env python3
"""
generate_activity_log.py

Generate Daily Activity Log Excel files (monthly or full-year).

USAGE
-----
Single month:
  python3 generate_activity_log.py -m 10 -y 2028
  python3 generate_activity_log.py -m 10 -y 2028 -o out_2028

Whole year (months 1..12):
  python3 generate_activity_log.py -y 2028 --all-months
  python3 generate_activity_log.py -y 2028 --all-months -o out_2028

Show help:
  python3 generate_activity_log.py -h

NOTES
-----
- Output filename format: Daily_Activity_Log_<YEAR>_<MM>.xlsx
- Week labels use Monday-start calendar week buckets:
  W1 contains the 1st of the month; W5/W6 may occur.
"""

from __future__ import annotations

import argparse
import calendar
import sys
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# -----------------------------
# Core logic
# -----------------------------

def week_label_in_month(d: date, week_starts_on: int = 0) -> str:
    """Return week-of-month label (W1, W2, ...) using calendar week buckets.

    Rule:
    - Weeks start on Monday by default (week_starts_on=0).
    - W1 is the week that contains the 1st day of the month (even if partial).
    - Each next week-start boundary increments the W number.
    - Some months will have W5 and occasionally W6.
    """
    first = d.replace(day=1)
    offset = (first.weekday() - week_starts_on) % 7  # 0..6
    day_index = d.day - 1  # 0-based
    wk = ((offset + day_index) // 7) + 1
    return f"W{wk}"


def build_activity_log(month: int, year: int, output_path: str | None = None) -> str:
    """Create an Excel Daily Activity Log for the given month/year.

    Args:
        month: 1..12
        year: 1900..9999
        output_path: Full output file path. If None, writes to CWD with default naming.

    Returns:
        Created file path as string.
    """
    if not (1 <= month <= 12):
        raise ValueError("month must be 1..12")
    if not (1900 <= year <= 9999):
        raise ValueError("year must be a 4-digit year")

    month_name = calendar.month_name[month]

    # Default output: current working directory
    out_file = Path(output_path) if output_path else Path(f"Daily_Activity_Log_{year}_{month:02d}.xlsx")
    out_file.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = f"{year}-{month:02d}"

    # -----------------
    # Styles
    # -----------------
    title_font = Font(bold=True, size=14, color="FFFFFF")
    header_font = Font(bold=True, color="FFFFFF")
    label_font = Font(bold=True)

    fill_title = PatternFill("solid", fgColor="1F4E79")
    fill_header = PatternFill("solid", fgColor="2F5597")
    fill_weekend = PatternFill("solid", fgColor="FFF2CC")  # light yellow
    fill_weekday = PatternFill("solid", fgColor="FFFFFF")

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left_top_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)

    thin = Side(style="thin", color="A6A6A6")
    border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

    # -----------------
    # Title row
    # -----------------
    ws.merge_cells("A1:D1")
    ws["A1"].value = f"Daily Activity Log ΓÇö {month_name} {year}"
    ws["A1"].font = title_font
    ws["A1"].fill = fill_title
    ws["A1"].alignment = align_center

    # Document the week rule so itΓÇÖs self-explanatory
    ws.merge_cells("A2:D2")
    ws["A2"].value = (
        "Week rule: Monday-start calendar week buckets; W1 contains the 1st of the month (W5/W6 may occur)."
    )
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Name / Position placeholders
    ws["A3"].value = "Name"
    ws["A3"].font = label_font
    ws["B3"].value = "<your name>"
    ws["C3"].value = "Position"
    ws["C3"].font = label_font
    ws["D3"].value = "<your position>"

    for cell in ("A3", "B3", "C3", "D3"):
        ws[cell].alignment = align_center
        ws[cell].border = border_all

    # -----------------
    # Table header
    # -----------------
    header_row = 5
    headers = ["Week", "Date", "Day", "Activity Log"]

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_all

    # -----------------
    # Fill daily rows
    # -----------------
    _, last_day = calendar.monthrange(year, month)
    row = header_row + 1

    for day in range(1, last_day + 1):
        d = date(year, month, day)
        is_weekend = d.weekday() >= 5  # Sat/Sun
        week_label = week_label_in_month(d, week_starts_on=0)  # Monday-start
        day_name = d.strftime("%a")  # Mon, Tue, ...

        ws.cell(row=row, column=1, value=week_label)
        date_cell = ws.cell(row=row, column=2, value=d)  # real Excel date
        ws.cell(row=row, column=3, value=day_name)
        ws.cell(row=row, column=4, value="")

        date_cell.number_format = "d mmmm yyyy"  # e.g., 2 October 2026

        fill = fill_weekend if is_weekend else fill_weekday

        for col in range(1, 5):
            cell = ws.cell(row=row, column=col)
            cell.border = border_all
            cell.fill = fill
            cell.alignment = align_left_top_wrap if col == 4 else align_center

        ws.row_dimensions[row].height = 30
        row += 1

    # -----------------
    # Column widths
    # -----------------
    widths = {
        1: 7,   # Week
        2: 18,  # Date
        3: 10,  # Day
        4: 70,  # Activity Log
    }
    for col_idx, w in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    # Freeze panes below header
    ws.freeze_panes = "A6"

    # Print setup (optional)
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    wb.save(out_file)
    return str(out_file)


# -----------------------------
# CLI / orchestration
# -----------------------------

@dataclass(frozen=True)
class RunPlan:
    year: int
    months: list[int]
    outdir: Path
    prefix: str = "Daily_Activity_Log"


def _validate_year(v: str) -> int:
    try:
        year = int(v)
    except ValueError as e:
        raise argparse.ArgumentTypeError("year must be an integer") from e
    if not (1900 <= year <= 9999):
        raise argparse.ArgumentTypeError("year must be between 1900 and 9999")
    return year


def _validate_month(v: str) -> int:
    try:
        month = int(v)
    except ValueError as e:
        raise argparse.ArgumentTypeError("month must be an integer") from e
    if not (1 <= month <= 12):
        raise argparse.ArgumentTypeError("month must be between 1 and 12")
    return month


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate an Excel Daily Activity Log for a given month/year or a whole year.",
        formatter_class=argparse.RawTextHelpFormatter,
        epilog=(
            "Examples:\n"
            "  python3 generate_activity_log.py -m 10 -y 2028\n"
            "  python3 generate_activity_log.py -y 2028 --all-months\n"
            "  python3 generate_activity_log.py -y 2028 --all-months -o out_2028\n"
        ),
    )

    parser.add_argument("-y", "--year", type=_validate_year, help="Year (e.g., 2028)")
    group = parser.add_mutually_exclusive_group()
    group.add_argument("-m", "--month", type=_validate_month, help="Month number (1-12)")
    group.add_argument("--all-months", action="store_true", help="Generate all months (1-12) for the given year")

    parser.add_argument(
        "-o", "--outdir",
        default=None,
        help=(
            "Output directory.\n"
            "Default: current directory for single-month; ./out_<YEAR> for --all-months"
        ),
    )
    parser.add_argument(
        "--prefix",
        default="Daily_Activity_Log",
        help="Output filename prefix (default: Daily_Activity_Log)",
    )

    return parser.parse_args(argv)


def make_plan(args: argparse.Namespace) -> RunPlan:
    # Backward-compatible prompts if args are missing
    year = args.year if args.year is not None else int(input("Enter year (e.g., 2028): ").strip())

    if args.all_months:
        months = list(range(1, 13))
        default_outdir = Path.cwd() / f"out_{year}"
    else:
        month = args.month if args.month is not None else int(input("Enter month (1-12): ").strip())
        if not (1 <= month <= 12):
            raise ValueError("month must be 1..12")
        months = [month]
        default_outdir = Path.cwd()

    outdir = Path(args.outdir).expanduser().resolve() if args.outdir else default_outdir.resolve()
    outdir.mkdir(parents=True, exist_ok=True)

    return RunPlan(year=year, months=months, outdir=outdir, prefix=args.prefix)


def main(argv: list[str]) -> int:
    args = parse_args(argv)

    # If user provided year but neither month nor --all-months, treat as "ask month"
    # (keeps interactive behavior)
    if args.year is not None and (args.month is None and not args.all_months):
        # fall back to prompt in make_plan()
        pass

    plan = make_plan(args)

    for m in plan.months:
        output_file = plan.outdir / f"{plan.prefix}_{plan.year}_{m:02d}.xlsx"
        created = build_activity_log(m, plan.year, output_path=str(output_file))
        print(f"Created: {created}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
